#!/usr/bin/env python3
"""Generate the KOZ280 field metadata from the official e-Tax Excel workbooks."""

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
STRUCTURE = ROOT / "e-taxall/09XML構造設計書等【所得税】/XML構造設計書(所得-申請)Ver21x.xlsx"
FIELDS = ROOT / "e-taxall/09XML構造設計書等【所得税】/帳票フィールド仕様書(所得-申請)Ver21x.xlsx"
OUTPUT = ROOT / "src/official-pko0420-field-spec.js"

COMPONENTS = {
    "yy": ["era", "yy"],
    "yymmdd": ["era", "yy", "mm", "dd"],
    "zipcode": ["zip1", "zip2"],
    "tel-number": ["tel1", "tel2", "tel3"],
    "kubun": ["kubun_CD"],
    "kubun2": ["kubun_CD"],
    "zeimusho": ["zeimusho_NM"],
}


def clean(value):
    return "" if value is None else str(value).strip()


def integer(value, default=1):
    try:
        return int(value)
    except (TypeError, ValueError):
        match = re.match(r"\s*(\d+)", clean(value))
        return int(match.group(1)) if match else default


def read_structure():
    sheet = load_workbook(STRUCTURE, read_only=True, data_only=True)["KOZ280"]
    stack = {}
    records = {}
    for row in list(sheet.iter_rows(values_only=True))[4:]:
        if not row[0] or not row[18]:
            continue
        level = integer(row[1], 0)
        if not level:
            continue
        tag = clean(row[18])
        stack[level] = tag
        stack = {key: value for key, value in stack.items() if key <= level}
        records[tag] = {
            "tag": tag,
            "path": "/".join(stack[key] for key in sorted(stack)),
            "commonType": clean(row[12]),
            "minOccurs": integer(row[13], 0),
            "maxOccurs": None if clean(row[14]) in ("", "unbounded") else integer(row[14]),
            "officialId": clean(row[15]),
            "idref": clean(row[16]),
        }
    return records


def read_fields(structure):
    sheet = load_workbook(FIELDS, read_only=True, data_only=True)["KOZ280"]
    grouped = defaultdict(list)
    for row in list(sheet.iter_rows(values_only=True))[3:]:
        if not row[0] or not row[12]:
            continue
        grouped[clean(row[12])].append(row)

    result = []
    for tag, rows in grouped.items():
        meta = structure[tag]
        component_names = COMPONENTS.get(meta["commonType"], [])
        repeat_tag = next((part for part in meta["path"].split("/") if (structure.get(part, {}).get("maxOccurs") or 1) > 1), "")
        repeats = max(integer(row[5], 1) for row in rows)
        for occurrence in range(1, repeats + 1):
            for index, row in enumerate(rows):
                item = integer(row[0])
                group_name = clean(row[3]) or clean(rows[0][3]) or tag
                name = clean(row[4]) or group_name
                label = f"{group_name} / {name}" if name != group_name else group_name
                if repeats > 1:
                    label = f"{label}（{occurrence}）"
                input_type = clean(row[1])
                component = component_names[index] if index < len(component_names) else ""
                result.append({
                    "id": f"pko_{item}_{occurrence}",
                    "item": item,
                    "label": label,
                    "group": group_name,
                    "name": name,
                    "page": 1,
                    "xmlPath": meta["path"],
                    "tag": tag,
                    "occurrence": occurrence,
                    "repeats": repeats,
                    "repeatTag": repeat_tag,
                    "component": component,
                    "inputType": input_type,
                    "type": "number" if input_type == "数値" else "text",
                    "format": clean(row[6]),
                    "inputCheck": clean(row[7]),
                    "range": clean(row[9]),
                    "calculation": clean(row[8]),
                    "calculationNo": clean(row[10]),
                    "note": clean(row[11]),
                    "officialId": meta["officialId"],
                    "idref": meta["idref"],
                    "commonType": meta["commonType"],
                    "minOccurs": meta["minOccurs"],
                    "maxOccurs": meta["maxOccurs"],
                    "pathMeta": [
                        {
                            "tag": part,
                            "min": structure.get(part, {}).get("minOccurs", 0),
                            "max": structure.get(part, {}).get("maxOccurs"),
                            "type": structure.get(part, {}).get("commonType", ""),
                        }
                        for part in meta["path"].split("/")
                    ],
                })
    return sorted(result, key=lambda entry: (entry["item"], entry["occurrence"]))


def main():
    structure = read_structure()
    fields = read_fields(structure)
    tags = {field["tag"] for field in fields}
    payload = json.dumps(fields, ensure_ascii=False, separators=(",", ":"))
    source = (
        "// Generated from the official KOZ280 Ver21 XML structure and field specification workbooks.\n"
        f"export const PKO0420_FIELD_SPEC = {payload};\n"
        "export const PKO0420_SPEC_COUNTS = "
        + json.dumps({"sourceRows": 114, "fields": len(fields), "tags": len(tags), "pages": 1}, ensure_ascii=False)
        + ";\n"
    )
    OUTPUT.write_text(source, encoding="utf-8")
    print(f"wrote {OUTPUT}: {len(fields)} fields, {len(tags)} tags")


if __name__ == "__main__":
    main()
